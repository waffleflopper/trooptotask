"""Builds the PowerPoint version of the deck from deck-render.mjs's slide images.

    node deck-render.mjs && ~/.cache/t2t-pitch-audio/venv/bin/python deck-pptx.py
    → out/troop-to-task-arctic-ascent-slides.pptx

Each slide is the rendered 1920×1080 image. The opener slide embeds the final video; in PowerPoint the next
click on that slide starts it.
"""
import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches

HERE = Path(__file__).parent
OUT = HERE / 'out'
SLIDES = OUT / 'slides'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]

for s in json.loads((SLIDES / 'slides.json').read_text()):
    slide = prs.slides.add_slide(blank)
    if s['video']:
        slide.shapes.add_movie(
            str(OUT / 'troop-to-task-arctic-ascent-final.mp4'), Emu(0), Emu(0), W, H,
            poster_frame_image=str(OUT / 'poster.png'), mime_type='video/mp4',
        )
    else:
        slide.shapes.add_picture(str(SLIDES / s['file']), Emu(0), Emu(0), W, H)

dest = OUT / 'troop-to-task-arctic-ascent-slides.pptx'
prs.save(dest)
print(f'wrote {dest.relative_to(HERE)}')
